import io
import re
import unicodedata
from dataclasses import dataclass, field

import pypdf
from docx import Document

# Fase 1 del agente de transformacion (ver agente/CONTRATO-AGENTE.md).
# Extraccion de texto sin LLM: si esto falla o detecta NO_TEXT_LAYER, nunca se llama a Claude.

MIN_CHARS = 400
MIN_WORDS = 40
MAX_CHARS = 20000

# Empiricamente derivado del PDF real de Shirley Mercado extraido con pypdf: el kerning de
# ciertas fuentes hace que headers como "PROFILE"/"EXPERIENCE"/"EDUCATION"/"SKILLS"/
# "CERTIFICATIONS"/"REFERENCES" salgan partidos letra por letra ("PR O FIL E", "E XPER IENC E").
# Heuristica: una linea completa de >=3 tokens, cada uno 1-4 letras mayusculas, se colapsa sin
# espacios. No es perfecto (una linea real tipo "SAN JUAN PR" tambien matchearia) -- aceptable
# porque esta normalizacion solo afecta el texto que ve el LLM, nunca el grounding contra el
# texto crudo.
_LETTER_SPACING_RE = re.compile(r"^(?:[A-Z]{1,4}\s+){2,}[A-Z]{1,4}$")

# Empiricamente derivado del PDF real de Ibrahim Rivas Andino (4 sep 2026): templates de
# 2 columnas ponen la fecha de cada rol en un bloque de texto flotante separado del
# titulo/empresa. pypdf extrae en el orden del stream interno del PDF, no en orden visual,
# asi que estas fechas terminan agrupadas en un punto random del texto (a veces todas juntas
# al principio del documento), totalmente desconectadas del rol al que describen -- el LLM
# termina adivinando mal a que rol pertenece cada fecha (o inventando precision que no esta
# en la fuente). grounding.py no lo bloquea porque no es una invencion de contenido, es una
# desconexion posicional.
_PDF_DATE_FRAGMENT_RE = re.compile(
    r"^(?:[A-Za-z]{3,9}\.?\s+)?\d{4}\s*"
    r"(?:[-–—]\s*(?:(?:[A-Za-z]{3,9}\.?\s+)?\d{4}|Present|Actualidad|Actual|Hoy|A la fecha))?\.?$",
    re.IGNORECASE,
)
_DATE_ANCHOR_Y_TOLERANCE = 20  # puntos PDF (~una linea de texto)
_MIN_ANCHOR_LEN = 4  # descarta viñetas/iconos/letras sueltas como ancla (encontrado con
# Yajaira Ortiz Ruiz y Luis Antonio Garcia Sanchez: sin este piso, un glifo de telefono o
# una letra suelta partida por kerning le ganaba por y a la linea real del rol/empresa)
_MAX_NEWLINES_ALREADY_NEAR = 1  # a lo sumo 1 salto de linea entre fecha y ancla se considera
# ya bien ubicado (ej. Ruth Sotomayor: compania y fecha en 2 lineas consecutivas -- ambiguo
# para un humano pero no para el LLM). Nunca una ventana de caracteres: un padding con muchos
# espacios literales para alinear a la derecha (visto en el PDF real de Luis Antonio Garcia
# Sanchez) infla la distancia en caracteres de un caso ya correcto y lo hace parecer roto.
_PLAUSIBLE_YEAR_RANGE = range(1950, 2036)  # descarta un numero de 4 digitos que no es un
# anio real (encontrado con Yajaira Ortiz Ruiz: "0851", los ultimos 4 digitos de un telefono
# partido por kerning en fragmentos separados, matcheaba el regex de fecha sin este chequeo)


def _is_plausible_date_fragment(text):
    if not _PDF_DATE_FRAGMENT_RE.match(text):
        return False
    return all(int(y) in _PLAUSIBLE_YEAR_RANGE for y in re.findall(r"\d{4}", text))


class ExtractionError(Exception):
    def __init__(self, code, detail):
        self.code = code
        self.detail = detail
        super().__init__(f"{code}: {detail}")


@dataclass
class ExtractResult:
    text: str
    kind: str
    units: int
    truncated: bool = False
    warnings: list = field(default_factory=list)


def _sniff_kind(data: bytes) -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK":
        return "docx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "legacy_doc"
    return "text"


def _page_text_fragments(page):
    """(texto, x, y) de cada fragmento de texto de la pagina, en coordenadas absolutas
    (cm compuesto con tm -- tm solo no alcanza cuando el fragmento vive dentro de un
    Form XObject con su propio cm, como el text-box flotante de fechas de Ibrahim)."""
    fragments = []

    def visitor(text, cm, tm, font_dict, font_size):
        stripped = text.strip()
        if not stripped:
            return
        a, b, c, d, e, f = tm
        A, B, C, D, E, F = cm
        x = A * e + C * f + E
        y = B * e + D * f + F
        fragments.append((stripped, x, y))

    page.extract_text(visitor_text=visitor)
    return fragments


def _find_nth_occurrence(haystack, needle, n):
    """Indice de la ocurrencia n-esima (0-indexed) de needle en haystack, o -1. Necesario
    porque un mismo valor de fecha (ej. 'May 2024') puede repetirse como fecha de inicio
    de un rol y fecha de fin de otro -- text.find() simple siempre agarra la primera
    ocurrencia, sin importar cual de los fragmentos reales se esta procesando."""
    start = 0
    idx = -1
    for _ in range(n + 1):
        idx = haystack.find(needle, start)
        if idx == -1:
            return -1
        start = idx + 1
    return idx


# 'MARCH 2006- JANUARY 2013' es un fragmento valido de por si, pero contiene '2006' como
# substring exacto -- si el propio anio '2006' aparece ademas como fecha suelta en otra
# parte del CV (real: Yajaira Ortiz Ruiz, seccion Education), _find_nth_occurrence lo
# encuentra "adentro" del fragmento mas largo por error. Solo aplica a fechas que empiezan
# con un digito (un anio suelto o "YYYY - ...") -- un fragmento como "March 2026 – Present"
# ya es especifico de sobra y no tiene este riesgo de colision.
_MONTH_WORD_BEFORE_RE = re.compile(r"[A-Za-z]{3,9}\.?\s*$")


def _find_nth_date_occurrence(haystack, needle, n):
    if not needle[0].isdigit():
        return _find_nth_occurrence(haystack, needle, n)
    count = 0
    start = 0
    while True:
        idx = haystack.find(needle, start)
        if idx == -1:
            return -1
        prefix = haystack[max(0, idx - 12):idx]
        if not _MONTH_WORD_BEFORE_RE.search(prefix):
            if count == n:
                return idx
            count += 1
        start = idx + 1


def _reattach_floating_dates(page, text):
    """Corrige fechas de rol desconectadas de su titulo/empresa por un template de 2
    columnas (ver comentario de _PDF_DATE_FRAGMENT_RE). Deliberadamente conservador:
    solo se toca una fecha si su vecino mas cercano por posicion (y) NO esta ya a lo sumo
    a _MAX_NEWLINES_ALREADY_NEAR saltos de linea de distancia en el texto por defecto.
    Confirmado sin falsos positivos contra 4 CVs reales -- Shirley, Yajaira, Luis, Ruth --
    cuyas fechas tambien vienen como fragmento aparte pero el orden del stream ya las deja
    razonablemente cerca de su rol.

    El rank (cuantas veces aparecio ya este mismo valor de texto en orden de fragmentos)
    se calcula una sola vez contra las posiciones originales -- nunca cambia, aunque el
    texto se edite despues -- porque una reubicacion mueve una ocurrencia existente, no
    agrega ni quita ocurrencias del valor."""
    fragments = _page_text_fragments(page)
    if not fragments:
        return text

    relocations = []
    seen_counts = {}
    for i, (frag_text, frag_x, frag_y) in enumerate(fragments):
        rank = seen_counts.get(frag_text, 0)
        seen_counts[frag_text] = rank + 1
        if not _is_plausible_date_fragment(frag_text):
            continue

        anchor_text, anchor_rank, best_dy = None, None, None
        anchor_seen_counts = {}
        for j, (other_text, other_x, other_y) in enumerate(fragments):
            other_rank = anchor_seen_counts.get(other_text, 0)
            anchor_seen_counts[other_text] = other_rank + 1
            if j == i or len(other_text) < _MIN_ANCHOR_LEN or _is_plausible_date_fragment(other_text):
                continue
            dy = abs(other_y - frag_y)
            if dy <= _DATE_ANCHOR_Y_TOLERANCE and (best_dy is None or dy < best_dy):
                anchor_text, anchor_rank, best_dy = other_text, other_rank, dy
        if anchor_text is not None:
            relocations.append((frag_text, rank, anchor_text, anchor_rank))

    for date_text, date_rank, anchor_text, anchor_rank in relocations:
        date_idx = _find_nth_date_occurrence(text, date_text, date_rank)
        anchor_idx = _find_nth_occurrence(text, anchor_text, anchor_rank)
        if date_idx == -1 or anchor_idx == -1:
            continue
        lo, hi = sorted((date_idx, anchor_idx))
        if text.count("\n", lo, hi) <= _MAX_NEWLINES_ALREADY_NEAR:
            continue  # ya estan razonablemente cerca en el texto extraido -- no tocar

        text = text[:date_idx] + text[date_idx + len(date_text):]
        anchor_idx = _find_nth_occurrence(text, anchor_text, anchor_rank)
        line_end = text.find("\n", anchor_idx)
        insert_at = len(text) if line_end == -1 else line_end
        text = text[:insert_at] + f"\t{date_text}" + text[insert_at:]

    return text


def _extract_pdf(data: bytes):
    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        page_text = _reattach_floating_dates(page, page_text)
        pages.append(page_text)
    return "\n".join(pages), len(reader.pages)


def _extract_docx(data: bytes):
    doc = Document(io.BytesIO(data))
    header_footer_paragraphs = []
    for section in doc.sections:
        header_footer_paragraphs += [p.text for p in section.header.paragraphs if p.text.strip()]
        header_footer_paragraphs += [p.text for p in section.footer.paragraphs if p.text.strip()]
    table_rows = [
        "\t".join(cell.text for cell in row.cells)
        for table in doc.tables
        for row in table.rows
    ]
    units = header_footer_paragraphs + [p.text for p in doc.paragraphs] + table_rows
    return "\n".join(units), len(units)


def _fix_letter_spacing(text: str) -> str:
    fixed_lines = []
    for line in text.split("\n"):
        stripped = line.strip()
        if _LETTER_SPACING_RE.match(stripped):
            fixed_lines.append(stripped.replace(" ", ""))
        else:
            fixed_lines.append(line)
    return "\n".join(fixed_lines)


def _sanitize(text: str) -> str:
    text = "".join(
        ch for ch in text
        if ch in ("\n", "\t") or unicodedata.category(ch)[0] != "C"
    )
    text = unicodedata.normalize("NFKC", text)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    cleaned, blank_run = [], 0
    for line in lines:
        if line == "":
            blank_run += 1
            if blank_run > 1:
                continue
        else:
            blank_run = 0
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def extract_text(filename: str, data: bytes) -> ExtractResult:
    kind = _sniff_kind(data)

    if kind == "legacy_doc":
        raise ExtractionError(
            "CORRUPT_FILE", f"{filename}: formato .doc legacy (binario, pre-2007) no soportado, se requiere PDF o DOCX"
        )

    try:
        if kind == "pdf":
            raw, units = _extract_pdf(data)
        elif kind == "docx":
            raw, units = _extract_docx(data)
        else:
            kind = "text"
            raw, units = data.decode("utf-8", errors="replace"), 1
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            "CORRUPT_FILE", f"no se pudo leer {filename} como {kind}: {exc}"
        ) from exc

    raw = _fix_letter_spacing(raw)
    clean = _sanitize(raw)

    words = clean.split()
    if len(clean) < MIN_CHARS or len(words) < MIN_WORDS:
        raise ExtractionError(
            "NO_TEXT_LAYER",
            f"solo {len(clean)} caracteres / {len(words)} palabras extraidas de "
            f"{units} pagina(s)/parrafo(s) ({filename}) -- CV probablemente escaneado "
            "o sin capa de texto",
        )

    warnings = []
    truncated = False
    if len(clean) > MAX_CHARS:
        clean = clean[:MAX_CHARS]
        truncated = True
        warnings.append(f"SOURCE_TRUNCATED: texto cortado a {MAX_CHARS} caracteres")

    return ExtractResult(text=clean, kind=kind, units=units, truncated=truncated, warnings=warnings)
