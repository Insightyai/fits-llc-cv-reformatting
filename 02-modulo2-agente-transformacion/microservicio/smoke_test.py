import io
import json
from pathlib import Path

import docx
from docxtpl import DocxTemplate

from blocks import CONTEXT_BUILDERS, strip_empty_paragraphs

BASE_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = BASE_DIR / "templates" / "anotados"
FIXTURES_DIR = BASE_DIR / "contrato-datos" / "fixtures"

TEMPLATES = {
    "new_format": TEMPLATES_DIR / "New Format Resume Template.docx",
    "non_template": TEMPLATES_DIR / "Non Template Resume.docx",
    "bd_format": TEMPLATES_DIR / "BD - Resume Template.docx",
}

FIXTURES = {
    "shirley-mercado": FIXTURES_DIR / "shirley-mercado.json",
    "adversarial": FIXTURES_DIR / "adversarial.json",
}


def render(template_id, cv):
    tpl = DocxTemplate(str(TEMPLATES[template_id]))
    context = CONTEXT_BUILDERS[template_id](cv)
    tpl.render(context, autoescape=True)
    buf = io.BytesIO()
    tpl.save(buf)
    buf.seek(0)
    rendered_doc = docx.Document(buf)
    strip_empty_paragraphs(rendered_doc)
    buf = io.BytesIO()
    rendered_doc.save(buf)
    buf.seek(0)
    return buf.read()


def check(template_id, fixture_name, cv):
    raw = render(template_id, cv)
    d = docx.Document(io.BytesIO(raw))
    full_text = "\n".join(p.text for p in d.paragraphs)

    assert "{{" not in full_text and "{%" not in full_text, "quedo un tag sin resolver"
    assert "None" not in full_text, "un campo null se imprimio como 'None' literal"
    assert cv["full_name"].lower() in full_text.lower(), "full_name no aparece en el render"
    assert "++" not in full_text, "years_experience duplico el '+' (bug '8++ YRS. OF EXP.')"
    assert "+ YRS. OF EXP." not in full_text or cv.get("years_experience"), (
        "quedo un '+ YRS. OF EXP.' huerfano con years_experience null"
    )

    for exp in cv["experience"]:
        for b in exp["bullets"]:
            assert b in full_text, f"bullet perdido: {b!r}"

    if template_id == "bd_format":
        assert "Resides in" not in full_text, "BD Format ya no debe mostrar 'Resides in ...'"
        # BD Format no renderiza skills_items (ver blocks.build_bd_format_context,
        # cubierto por test_blocks.py) -- no se repite ese chequeo aca porque algunos
        # skills adversariales (ej. "R&D") coinciden a proposito con texto del summary
        # y de los bullets, y darian un falso positivo.
    else:
        for s in cv["skills"]:
            assert s in full_text, f"skill perdida: {s!r}"

    print(f"  OK  {template_id:15s} + {fixture_name:16s} -> {len(raw)} bytes, sin tags sueltos, sin 'None', bullets intactos")


for fixture_name, path in FIXTURES.items():
    cv = json.load(open(path, encoding="utf-8"))
    print(f"=== fixture: {fixture_name} ===")
    for template_id in TEMPLATES:
        check(template_id, fixture_name, cv)

print("\nSMOKE TEST OK — 3 templates x 2 fixtures, 6/6 combinaciones pasaron")
